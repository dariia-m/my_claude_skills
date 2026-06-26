# -*- coding: utf-8 -*-
"""build_xlsx.py - turn a curated JSON list of paper records into a formatted .xlsx
literature table.

Usage:
  python build_xlsx.py papers.json -o literature_<topic>.xlsx

Standard columns (in order): authors, year, title, journal, full_citation, keywords,
abstract, citation_count, paper_type, oa_status, link. Any EXTRA key in a record (other
than the internal helper keys) becomes an additional column - that is how per-run custom
columns work: add the key to each record before calling this. `full_citation` is built in
AEA style from the fields when a record doesn't already provide one.
"""
import sys, json, argparse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

STANDARD = ["authors", "year", "title", "journal", "full_citation", "keywords",
            "abstract", "citation_count", "paper_type", "oa_status", "link"]
INTERNAL = {"doi", "source_type", "type", "volume", "issue", "pages", "is_oa", "pdf_candidates",
            "source", "fields", "external"}  # provenance/helper keys, never table columns
WIDTHS = {"authors": 26, "year": 6, "title": 46, "journal": 26, "full_citation": 52,
          "keywords": 30, "abstract": 62, "citation_count": 9, "paper_type": 18,
          "oa_status": 11, "link": 42}
WRAP = {"title", "full_citation", "keywords", "abstract", "authors"}

def aea_authors(names):
    def split(n):
        parts = n.strip().split()
        return (" ".join(parts[:-1]), parts[-1]) if len(parts) > 1 else ("", n)
    sp = [split(n) for n in names]
    if not sp:
        return ""
    if len(sp) == 1:
        g, s = sp[0]; return f"{s}, {g}".strip().rstrip(",")
    head = f"{sp[0][1]}, {sp[0][0]}".strip()
    mids = [f"{g} {s}".strip() for g, s in sp[1:-1]]
    last = f"{sp[-1][0]} {sp[-1][1]}".strip()
    return ", ".join([head] + mids + [f"and {last}"]) if mids else f"{head}, and {last}"

def aea_citation(r):
    auth = aea_authors(r.get("authors") or [])
    yr = r.get("year") or "n.d."
    title = (r.get("title") or "").rstrip(".")
    jr = r.get("journal") or ""
    auth = (auth + ".") if auth and not auth.endswith(".") else auth
    cite = f'{auth} {yr}. "{title}." {jr}'.strip()
    if r.get("volume"):
        cite += f" {r['volume']}"
        if r.get("issue"):
            cite += f" ({r['issue']})"
    if r.get("pages"):
        cite += f": {r['pages']}"
    return cite.rstrip() + "."

def cellval(r, col):
    v = r.get(col)
    if col == "authors" and isinstance(v, list):
        return "; ".join(v)
    if col == "keywords" and isinstance(v, list):
        return "; ".join(v)
    if col == "full_citation" and not v:
        return aea_citation(r)
    return v

def build(records, out_path):
    custom = []
    for r in records:
        for k in r:
            if k not in STANDARD and k not in INTERNAL and k not in custom:
                custom.append(k)
    cols = [c for c in STANDARD if any(c in r or c == "full_citation" for r in records)] + custom

    wb = Workbook(); ws = wb.active; ws.title = "Literature"
    labels = {"oa_status": "OA Status", "full_citation": "Full Citation", "doi": "DOI",
              "citation_count": "Citation Count", "paper_type": "Paper Type"}
    hdr_fill = PatternFill("solid", fgColor="4C72B0"); hdr_font = Font(bold=True, color="FFFFFF")
    for j, col in enumerate(cols, 1):
        c = ws.cell(1, j, labels.get(col, col.replace("_", " ").title()))
        c.fill = hdr_fill; c.font = hdr_font; c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(j)].width = WIDTHS.get(col, 22)
    for i, r in enumerate(records, 2):
        for j, col in enumerate(cols, 1):
            val = cellval(r, col)
            cell = ws.cell(i, j, val if val is not None else "")
            cell.alignment = Alignment(wrap_text=(col in WRAP), vertical="top")
            if col == "link" and val:
                cell.hyperlink = val; cell.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(records)+1}"
    wb.save(out_path)
    return cols

if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("json_file")
    ap.add_argument("-o", "--out", required=True)
    a = ap.parse_args()
    recs = json.load(open(a.json_file, encoding="utf-8"))
    cols = build(recs, a.out)
    sys.stderr.write(f"Wrote {len(recs)} papers x {len(cols)} cols -> {a.out}\n  columns: {cols}\n")
